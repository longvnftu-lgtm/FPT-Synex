# -*- coding: utf-8 -*-
"""
Script CHUNG tao file Excel import cho cac skill base-* — dong bang format + tu verify.

Moi skill import/update trong bo chia se kem mot ban cua script nay trong thu muc scripts/.
- profile "crm":     cac skill CRM + base-dataset-2-import-field/-record + base-service-1-import-record
- profile "service": base-service-1-import-custom-field + base-schedule-1-import-phanca-daydu
- profile "generic": base-schedule-1-import-camau / -phanca-camau, base-checkin-1-import,
                     base-wework-1-import-field / -tasks, base-table-1-import-field / -record, base-test-1-import
- profile "payroll": bang luong (kem freeze/meta_cols/gray_zero)

Usage:
    python write_import_excel.py --spec "duong/dan/spec.json"

Spec JSON (UTF-8, ho tro tieng Viet day du — LUON ghi spec bang UTF-8):
{
  "output":        "D:/.../Import - CRM Pipeline.xlsx",
  "sheet_name":    "Exported file",
  "profile":       "crm",
  "headers":       ["Field name *", "Field key", "Field type *", ...],
  "rows":          [["Tên khách hàng", null, "text", "internal", ...], ...],
  "column_widths": [28, 18, 20, 14, 36, 40, 16, 24, 10, 24, 24, 22]
}
- "column_widths" tuy chon — bo qua thi auto theo noi dung (tran 60).
- Gia tri "" hoac null -> cell rong that su (None), khong phai chuoi rong.
- Gia tri so (int/float) giu nguyen kieu so — khong ep ve text (vd Deal value phai la so).

CHE DO NHIEU SHEET (khoa tuy chon "sheets") — dung cho file import 1 phieu kem
du lieu Bang cua Base Service 1.0:
{
  "output":  "D:/.../Import - QT-02 Phieu 1.xlsx",
  "profile": "crm",
  "sheets": [
    {"sheet_name": "Ticket Information",              "headers": [...], "rows": [[...]]},
    {"sheet_name": "service_thanh_phan_tham_du_phon", "headers": [...], "rows": [[...]]},
    {"sheet_name": "Worksheet",                       "headers": [],    "rows": []}
  ]
}
- Khong co "sheets" -> tu goi {sheet_name, headers, rows} cap cao thanh list 1 phan tu.
  Moi spec 1-sheet cu chay y nguyen, khong phai sua gi.
- Tuy chon (profile, column_widths, freeze, meta_cols, gray_zero, bold_rows,
  bold_rows_fill) tra o SHEET truoc, thieu thi lay o CAP WORKBOOK.
- Sheet RONG hoan toan: "headers": [] VA "rows": [] — khong ghi header, khong
  freeze, khong dat do rong cot.

CANH BAO — THU TU SHEET LA CO CHE LIEN KET DUY NHAT:
Base Service khop sheet Bang voi cot Bang tren sheet dau bang THU TU SHEET, khong
co cot ID nao lien ket. Doi thu tu hoac doi ten sheet = du lieu vao SAI BANG ma
khong co canh bao nao. Nen ten + thu tu sheet phai chep nguyen van tu file
template that cua quy trinh dich.

Quy tac dong bang:
- KHONG BAO GIO co dong nao sau dong du lieu cuoi (dong thua lam Base tu choi import).
- Khong dong [Example], khong dong ghi chu — canh bao gi thi in ra chat, khong nhet vao file.

Verify tich hop (tu chay sau khi ghi):
- Mo lai file, doi chieu: ten + THU TU sheet, header khop 100%, tong so dong = 1 + len(rows).
- In bang dem o-co-du-lieu theo tung cot de doi chieu voi danh sach da chot.
- Lech bat ky diem nao -> exit code 1 + bao loi ro. Exit 0 = file dung cau truc.
"""

import argparse
import json
import os
import re
import sys
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PROFILES = {
    # CRM Base 2.0 (field-import, record-import, update-record)
    "crm": {
        "header_fill": "1F3864",
        "header_font_color": "FFFFFF",
        "font_name": "Calibri",
        "font_size": 11,
        "align_h": "right",
        "wrap": True,
        "header_height": 30,
        "alt_fills": ("FFFFFF", "EBF0FA"),
        "border": False,
    },
    # Base Service 1.0 (giu tuong thich voi create_excel.py cu)
    "service": {
        "header_fill": "333333",
        "header_font_color": "FFFFFF",
        "font_name": "Calibri",
        "font_size": 10,
        "align_h": "left",
        "wrap": True,
        "header_height": None,
        "alt_fills": ("FFFFFF", "F7F7F7"),
        "border": True,
    },
    # Khong thuoc ho format nao cu the
    "generic": {
        "header_fill": "1B3A2D",
        "header_font_color": "FFFFFF",
        "font_name": "Calibri",
        "font_size": 11,
        "align_h": "left",
        "wrap": True,
        "header_height": None,
        "alt_fills": ("FFFFFF", "F8F8F8"),
        "border": False,
    },
    # Base Payroll 1.0 import (meta cols + zero xam — bat qua options)
    "payroll": {
        "header_fill": "1B3A2D",
        "header_font_color": "FFFFFF",
        "font_name": "Calibri",
        "font_size": 10,
        "align_h": "left",
        "wrap": True,
        "header_height": 45,
        "alt_fills": ("FFFFFF", "F2F7F4"),
        "border": True,
    },
}

# --- Profile sinh tu design system (them 08/2026) -------------------------------
# Ngoai 4 profile dong bang o tren, spec co the dung "profile": "ds:<slug>" de lay
# mau/font tu mot bo design system. Bo design system KHONG kem trong ban chia se nay:
# muon dung thi dat thu muc "design-systems/" + "ds_scripts/tokens_lib.py" canh file nay.
DS_SCRIPTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ds_scripts")
DS_PREFIX = "ds:"


def _mix(c1, c2, t):
    a = tuple(int(c1[i:i + 2], 16) for i in (0, 2, 4))
    b = tuple(int(c2[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(a[i] + (b[i] - a[i]) * t):02X}" for i in range(3))


def resolve_profile(name):
    """Tra ve dict style. Ten thuong -> PROFILES; 'ds:<slug>' -> dung tu design system."""
    if not str(name).startswith(DS_PREFIX):
        return PROFILES[name]

    slug = str(name)[len(DS_PREFIX):]
    sys.path.insert(0, DS_SCRIPTS)
    try:
        from tokens_lib import load_plain, type_scale
    except ImportError as e:
        sys.exit(f"LOI: khong nap duoc tokens_lib tu {DS_SCRIPTS} ({e}).")

    path = os.path.join(os.path.dirname(DS_SCRIPTS), "design-systems", slug,
                        "design-system.tokens.json")
    if not os.path.exists(path):
        sys.exit(f"LOI SPEC: khong co design system '{slug}' ({path}).")

    ds = load_plain(path)
    color, ext = ds["color"], ds.get("ext", {})
    table = ext.get("table", {})

    use_accent = table.get("header") == "accent"
    header_fill = (color["accent"] if use_accent else color["ink"]).lstrip("#")
    header_font = (color["accent-ink"] if use_accent else color["surface"]).lstrip("#")
    ground = color["ground"].lstrip("#")
    tint = float(table.get("zebra_tint", 0.11)) if table.get("zebra", True) else 0.0
    alt = _mix(ground, color["accent"].lstrip("#"), tint)

    body = ds["type"]["body"]["fontFamily"]
    return {
        "header_fill": header_fill,
        "header_font_color": header_font,
        "font_name": body[0] if isinstance(body, list) else body,
        "font_size": type_scale(ds)["body"],
        "align_h": "left",
        "wrap": True,
        "header_height": {"compact": None, "normal": 30, "roomy": 42}.get(
            ext.get("density", "normal"), 30),
        "alt_fills": (ground, alt),
        "border": table.get("borders", "horizontal") == "all",
    }


def valid_profile(name):
    return str(name).startswith(DS_PREFIX) or name in PROFILES


# Tuy chon them trong spec (ap dung moi profile):
#   "freeze":    "D2"   — o freeze panes, mac dinh "A2"
#   "meta_cols": 3      — N cot dau la cot meta: header fill 2E6B50, data fill E8F4F0
#   "gray_zero": true   — gia tri so 0 to font xam AAAAAA (de nhan ra o dien-mac-dinh)
#   "bold_rows": [0, 7] — index (0-based, theo rows[]) cua dong can in DAM.
#                         Dung cho Wework: dong "Nhom cong viec" phai in dam + ten ket thuc dau ":"
#                         thi Base moi nhan dien la nhom. Bo qua -> khong dong nao in dam (nhu cu).
#   "bold_rows_fill": "FFF2CC" — ma mau to nen cho cac dong trong bold_rows (tuy chon)
META_HEADER_FILL = "2E6B50"
META_DATA_FILL = "E8F4F0"
GRAY_ZERO_COLOR = "AAAAAA"

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def norm_cell(value):
    """'' / null -> None de cell rong that su; so giu nguyen kieu so."""
    if value is None:
        return None
    if isinstance(value, str) and value.strip() == "":
        return None
    return value


# Tuy chon co the khai o cap sheet, thieu thi lay o cap workbook.
INHERITED_KEYS = (
    "profile", "column_widths", "freeze", "meta_cols",
    "gray_zero", "bold_rows", "bold_rows_fill",
)
INVALID_SHEET_CHARS = set(r"[]:*?/\'")
MAX_SHEET_NAME = 31


def check_sheet_names(names):
    """Cong chan ten sheet — SAI LA DUNG, khong de openpyxl cat/doi ten ngam.

    Base khop sheet Bang voi cot Bang tren sheet dau bang THU TU sheet, khong co
    cot ID lien ket. Hai hanh vi ngam cua openpyxl deu tao loi IM LANG:
      - ten > 31 ky tu  -> chi warn(), file van ghi, Excel/Base cat khi doc.
      - ten trung nhau   -> avoid_duplicate_name() TU DOI TEN sheet sau.
    Ca hai deu cho ra file "tao thanh cong" nhung du lieu vao sai Bang.
    """
    errors = []
    for i, name in enumerate(names, 1):
        if not isinstance(name, str) or not name.strip():
            errors.append(f"Sheet #{i}: 'sheet_name' rong hoac khong phai chuoi.")
            continue
        if len(name) > MAX_SHEET_NAME:
            errors.append(
                f"Sheet #{i} '{name}': dai {len(name)} ky tu > {MAX_SHEET_NAME} (gioi han Excel). "
                "Phai ghi dung ten DA BI CAT trong file template goc, khong ghi ten day du."
            )
        bad = "".join(sorted(set(name) & INVALID_SHEET_CHARS))
        if bad:
            errors.append(f"Sheet #{i} '{name}': chua ky tu cam {bad!r}.")

    seen = {}
    for i, name in enumerate(names, 1):
        # Excel khong phan biet hoa/thuong khi so ten sheet
        key = str(name)[:MAX_SHEET_NAME].strip().lower()
        if key in seen:
            errors.append(
                f"Sheet #{i} '{name}' TRUNG TEN voi sheet #{seen[key]} sau khi cat {MAX_SHEET_NAME} ky tu. "
                "openpyxl se tu them so vao duoi -> lech ten/thu tu so voi template goc."
            )
        else:
            seen[key] = i

    if errors:
        print("LOI SPEC — ten sheet khong hop le:")
        for e in errors:
            print(f"  - {e}")
        sys.exit(1)


def normalize_sheets(spec):
    """Gom spec ve list sheet. Spec 1-sheet cu (khong co 'sheets') chay nguyen ven."""
    if "sheets" in spec:
        raw = spec["sheets"]
        if not isinstance(raw, list) or not raw:
            sys.exit("LOI SPEC: 'sheets' phai la list co it nhat 1 phan tu.")
    else:
        raw = [{
            "sheet_name": spec["sheet_name"],
            "headers": spec["headers"],
            "rows": spec["rows"],
        }]

    sheets = []
    for i, sh in enumerate(raw, 1):
        if not isinstance(sh, dict):
            sys.exit(f"LOI SPEC: sheet #{i} phai la object, dang la {type(sh).__name__}.")
        for key in ("sheet_name", "headers", "rows"):
            if key not in sh:
                sys.exit(f"LOI SPEC: sheet #{i} thieu khoa bat buoc '{key}'.")
        merged = dict(sh)
        for key in INHERITED_KEYS:  # thieu o sheet -> lay cap workbook
            if key not in merged and key in spec:
                merged[key] = spec[key]
        if not valid_profile(merged.get("profile", "generic")):
            sys.exit(
                f"LOI SPEC: sheet #{i} profile '{merged['profile']}' khong ton tai "
                f"(co: {list(PROFILES)}, hoac 'ds:<slug>')."
            )
        sheets.append(merged)

    check_sheet_names([s["sheet_name"] for s in sheets])
    return sheets


def build(spec):
    sheets = normalize_sheets(spec)

    # Chan loi cau truc truoc khi ghi
    for sh in sheets:
        for i, row in enumerate(sh["rows"]):
            if len(row) > len(sh["headers"]):
                sys.exit(
                    f"LOI SPEC: sheet '{sh['sheet_name']}' dong {i + 1} co "
                    f"{len(row)} gia tri > {len(sh['headers'])} cot header."
                )

    wb = Workbook()
    for idx, sh in enumerate(sheets):
        # Sheet dau PHAI dung wb.active — create_sheet cho ca sheet dau se de lai
        # sheet "Sheet" rong o vi tri 1 -> lech toan bo anh xa sheet Bang.
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = sh["sheet_name"]
        write_sheet(ws, sh)

    out = spec["output"]
    out_dir = os.path.dirname(out)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out)
    normalize_empty_cells(out)
    return out


def write_sheet(ws, sh):
    """Ghi 1 sheet. headers=[] va rows=[] -> sheet RONG hoan toan, khong cham gi ca."""
    headers = sh["headers"]
    rows = sh["rows"]

    # Thoat som PHAI dat truoc ca freeze_panes/column_dimensions — sheet rong ma
    # co <pane>/<cols> thi khong giong template goc nua.
    if not headers and not rows:
        return

    profile = resolve_profile(sh.get("profile", "generic"))

    header_font = Font(
        name=profile["font_name"], size=profile["font_size"],
        bold=True, color=profile["header_font_color"],
    )
    header_fill = PatternFill("solid", fgColor=profile["header_fill"])
    align = Alignment(horizontal=profile["align_h"], vertical="center", wrap_text=profile["wrap"])
    data_font = Font(name=profile["font_name"], size=profile["font_size"])

    meta_cols = int(sh.get("meta_cols", 0))
    gray_zero = bool(sh.get("gray_zero", False))
    meta_header_fill = PatternFill("solid", fgColor=META_HEADER_FILL)
    meta_data_fill = PatternFill("solid", fgColor=META_DATA_FILL)
    gray_font = Font(name=profile["font_name"], size=profile["font_size"], color=GRAY_ZERO_COLOR)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=profile["wrap"])

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = meta_header_fill if c <= meta_cols else header_fill
        cell.alignment = align
        if profile["border"]:
            cell.border = BORDER
    if profile["header_height"]:
        ws.row_dimensions[1].height = profile["header_height"]

    bold_rows = set(sh.get("bold_rows", []))
    bold_font = Font(name=profile["font_name"], size=profile["font_size"], bold=True)
    bold_fill_hex = sh.get("bold_rows_fill")
    bold_fill = PatternFill("solid", fgColor=bold_fill_hex) if bold_fill_hex else None

    for r_off, row in enumerate(rows):
        fill = PatternFill("solid", fgColor=profile["alt_fills"][r_off % 2])
        is_bold_row = r_off in bold_rows
        padded = list(row) + [None] * (len(headers) - len(row))
        for c, value in enumerate(padded, 1):
            v = norm_cell(value)
            cell = ws.cell(row=r_off + 2, column=c, value=v)
            is_num = isinstance(v, (int, float)) and not isinstance(v, bool)
            if is_bold_row:
                cell.font = bold_font
            else:
                cell.font = gray_font if (gray_zero and is_num and v == 0) else data_font
            if is_bold_row and bold_fill is not None:
                cell.fill = bold_fill
            else:
                cell.fill = meta_data_fill if c <= meta_cols else fill
            cell.alignment = align_right if is_num else align
            if profile["border"]:
                cell.border = BORDER

    widths = sh.get("column_widths")
    for c in range(1, len(headers) + 1):
        letter = ws.cell(row=1, column=c).column_letter
        if widths and c <= len(widths):
            ws.column_dimensions[letter].width = widths[c - 1]
        else:
            longest = max(
                [len(str(headers[c - 1]))]
                + [len(str(norm_cell(row[c - 1]) or "")) for row in rows if len(row) >= c]
            )
            ws.column_dimensions[letter].width = min(longest + 4, 60)

    ws.freeze_panes = sh.get("freeze", "A2")


def normalize_empty_cells(path):
    """Chuan hoa o TRONG ve dung dang Excel that ghi.

    Van de: o trong nhung CO dinh dang (mau nen, font) bi openpyxl ghi thanh
        <c r="B2" s="2" t="n"></c>
    tuc la khai bao KIEU SO (t="n") nhung khong co gia tri. Excel that thi ghi
        <c r="B2" s="2"/>
    khong co thuoc tinh t. Trinh import cua Base doc thay 't=n' + rong -> ep ve 0,
    nen moi o trong len he thong thanh so 0 (Nguoi giao viec = 0, Muc tieu = 0...).

    Ham nay go bo t="n" o cac cell rong sau khi openpyxl da ghi xong.
    Cell co gia tri (<v> hoac <is> ben trong) KHONG bi dung toi.
    """
    empty_cell = re.compile(r'<c([^>/]*?)\st="n"\s*(?:></c>|/>)')
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                xml = data.decode("utf-8")
                xml = empty_cell.sub(r"<c\1/>", xml)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


def verify_sheet(ws, sh):
    """Doi chieu 1 sheet voi spec cua no. Tra ve list loi (rong = dung)."""
    headers = sh["headers"]
    rows = sh["rows"]
    errors = []

    # Sheet RONG: openpyxl tra max_row=1/max_column=1 cho sheet khong co o nao,
    # nen KHONG duoc dung cong thuc 1+len(rows) o day (se ra 2 loi gia).
    if not headers and not rows:
        is_empty = (
            ws.max_row == 1 and ws.max_column == 1
            and ws.cell(row=1, column=1).value is None
        )
        if not is_empty:
            errors.append(
                f"[{ws.title}] Phai RONG hoan toan nhung dang co du lieu "
                f"({ws.max_row} dong x {ws.max_column} cot)."
            )
        print(f"Sheet: {ws.title} | RONG hoan toan (khong header, khong data) — dung template")
        return errors

    actual_headers = [ws.cell(row=1, column=c).value for c in range(1, len(headers) + 1)]
    if actual_headers != headers:
        for c, (exp, got) in enumerate(zip(headers, actual_headers), 1):
            if exp != got:
                errors.append(f"[{ws.title}] Header cot {c}: mong '{exp}' nhung file ghi '{got}'.")

    expected_rows = 1 + len(rows)
    if ws.max_row != expected_rows:
        errors.append(
            f"[{ws.title}] So dong lech: file co {ws.max_row}, mong {expected_rows} "
            f"(1 header + {len(rows)} data). Co the co DONG THUA sau data cuoi."
        )
    if ws.max_column > len(headers):
        errors.append(
            f"[{ws.title}] So cot lech: file co {ws.max_column}, mong {len(headers)}. "
            "Co the co COT THUA."
        )

    print(f"Sheet: {ws.title} | Dong: 1 header + {ws.max_row - 1} data | Cot: {ws.max_column}")
    print("Dem o co du lieu theo cot (doi chieu voi danh sach da chot):")
    for c, h in enumerate(headers, 1):
        n = sum(
            1 for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=c).value is not None
        )
        # Header co the chua xuong dong (vd cot "Danh muc phieu") — chi thay bang
        # dau cach KHI IN de bang khong vo cot; du lieu trong file giu nguyen van.
        label = " ".join(str(h).split())
        print(f"  {c:>2}. {label:<32}: {n}/{len(rows)}")
    return errors


def verify(spec):
    """Mo lai file vua ghi, doi chieu voi spec. Lech -> exit 1."""
    out = spec["output"]
    sheets = normalize_sheets(spec)
    errors = []

    wb = load_workbook(out, data_only=False)

    # THU TU sheet la co che lien ket duy nhat cua Base (khop sheet Bang voi cot
    # Bang tren sheet dau theo THU TU) -> phai khop tuyet doi, ke ca thu tu.
    expected_names = [s["sheet_name"] for s in sheets]
    if wb.sheetnames != expected_names:
        errors.append(
            "Ten/thu tu sheet LECH so voi spec (Base khop sheet Bang theo thu tu, "
            "lech = du lieu vao sai Bang):\n"
            f"      file: {wb.sheetnames}\n"
            f"      spec: {expected_names}"
        )

    print("=== VERIFY ===")
    print(f"File : {out}")
    for sh in sheets:
        name = sh["sheet_name"]
        if name not in wb.sheetnames:
            errors.append(f"Sheet '{name}' khong ton tai (co: {wb.sheetnames}).")
            continue
        errors.extend(verify_sheet(wb[name], sh))

    # O trong con khai kieu so (t="n") -> Base import se ep ve 0. Phai bang 0.
    # Quet TOAN BO xl/worksheets/*.xml -> dung cho ca workbook nhieu sheet.
    leftover = 0
    with zipfile.ZipFile(out, "r") as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                xml = z.read(name).decode("utf-8")
                leftover += len(re.findall(r'<c[^>/]*?\st="n"\s*(?:></c>|/>)', xml))
    print(f"O trong con khai kieu so (t=\"n\") — Base se doc thanh 0: {leftover} (phai = 0)")
    if leftover:
        errors.append(
            f"Con {leftover} o trong mang t=\"n\". Base import se hien thi thanh so 0. "
            "Kiem tra ham normalize_empty_cells()."
        )

    if errors:
        print("VERIFY THAT BAI:")
        for e in errors:
            print(f"  - {e}")
        sys.exit(1)
    print("VERIFY OK — cau truc file khop spec, khong dong/cot thua, o trong sach.")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--spec", required=True, help="Duong dan file spec JSON (UTF-8)")
    args = parser.parse_args()

    with open(args.spec, "r", encoding="utf-8-sig") as f:
        spec = json.load(f)

    if "output" not in spec:
        sys.exit("LOI SPEC: thieu khoa bat buoc 'output'.")
    if "sheets" not in spec:
        # Spec 1-sheet (mac dinh) — giu nguyen thong bao loi cu
        for key in ("sheet_name", "headers", "rows"):
            if key not in spec:
                sys.exit(f"LOI SPEC: thieu khoa bat buoc '{key}'.")
    if not valid_profile(spec.get("profile", "generic")):
        sys.exit(f"LOI SPEC: profile '{spec['profile']}' khong ton tai "
                 f"(co: {list(PROFILES)}, hoac 'ds:<slug>').")

    out = build(spec)
    print(f"Da tao file: {out}")
    verify(spec)


if __name__ == "__main__":
    main()
