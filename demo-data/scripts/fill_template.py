# -*- coding: utf-8 -*-
"""Do du lieu demo goc vao DUNG file mau record xuat tu tenant Base (giu nguyen thu tu cot).

Theo skill base-crm-account-import-record / base-crm-pipeline-import-record:
- Cot lay 100% tu file mau that (ten + thu tu), khong tu ghep cot.
- Cot nao trong file mau khong map duoc -> de trong + liet ke ra man hinh.
- Owner/Followers phai la username that -> lay tu personas.json, thieu thi DUNG.

Chay:
  python demo-data/scripts/fill_template.py --object accounts \
      --template "templated.file.report....xlsx" --personas demo-data/personas.json \
      --output "demo-data/Dữ liệu setup/CRM Sale/Account/Import - CRM Account record.xlsx"
  --object: accounts | contacts | deals_b2g | deals_xsell
"""
import argparse
import json
import os
import re
import subprocess
import sys
import unicodedata

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import demo_master as m  # noqa: E402

WRITER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "write_import_excel.py")
OWNER_KEYS = {"Account owner", "Deal owner", "Owner"}

# Ten cot thuong gap tren Base -> khoa trong du lieu goc (so sanh sau khi chuan hoa).
ALIASES = {
    "accounts": {
        "name": "Account name", "account name": "Account name", "ten cong ty": "Account name",
        "owner": "Account owner", "account owner": "Account owner",
        "value": "Account value", "account value": "Account value",
        "estimated lifetime value": "Account value", "currency": "Currency",
        "domain": "Domain", "address": "Address",
    },
    "contacts": {
        "name": "Contact name", "full name": "Contact name", "contact name": "Contact name",
        "title": "Title", "job title": "Title", "account": "Account", "account name": "Account",
        "company": "Account", "email": "Email", "phone": "Phone", "owner": "Owner",
        "contact owner": "Owner",
    },
    "deals": {
        "name": "Deal name", "deal name": "Deal name", "stage": "Stage", "value": "Deal value",
        "deal value": "Deal value", "currency": "Currency", "deal currency": "Currency",
        "account": "Account", "account name": "Account", "contact": "Contact",
        "contact name": "Contact", "owner": "Deal owner", "deal owner": "Deal owner",
        "probability": "Probability", "expected close": "Expected close date",
        "expected close date": "Expected close date", "overview": "Overview",
    },
}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"[\s_*]+", " ", s.lower()).strip()
    return s


def load_records(obj):
    data = m.build_all()
    return data[obj]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--object", required=True, choices=["accounts", "contacts", "deals_b2g", "deals_xsell"])
    ap.add_argument("--template", required=True, help="File mau record xuat tu Base cho dung object/pipeline")
    ap.add_argument("--personas", required=True, help="JSON {\"SALES_BAC\": \"username\", ...}")
    ap.add_argument("--output", required=True)
    ap.add_argument("--list-sep", default=";", help="Dau phan cach cho field nhieu lua chon (mac dinh ;)")
    a = ap.parse_args()

    wb = load_workbook(a.template, read_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    while headers and headers[-1] in (None, ""):
        headers.pop()

    with open(a.personas, encoding="utf-8") as fh:
        personas = json.load(fh)
    missing = [k for k in m.PERSONAS if not personas.get(k)]
    if missing:
        sys.exit(f"DUNG: personas.json chua co username that cho {missing}. Khong tu bia username.")

    records = load_records(a.object)
    aliases = ALIASES["deals" if a.object.startswith("deals") else a.object]
    keys_by_norm = {norm(k): k for k in records[0] if not k.startswith("_")}

    mapping, unmapped = [], []
    for h in headers:
        n = norm(h)
        key = keys_by_norm.get(n) or aliases.get(n)
        mapping.append(key)
        if key is None:
            unmapped.append(h)
    used = {k for k in mapping if k}
    unused = [k for k in keys_by_norm.values() if k not in used]

    rows = []
    for r in records:
        row = []
        for key in mapping:
            v = r.get(key) if key else None
            if key in OWNER_KEYS and v:
                v = personas[v]
            if isinstance(v, list):
                v = a.list_sep.join(v) if v else None
            row.append(v)
        rows.append(row)

    spec = {"output": a.output, "sheet_name": ws.title, "profile": "crm", "headers": headers, "rows": rows}
    spec_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_specs", f"fill_{a.object}.json")
    os.makedirs(os.path.dirname(spec_path), exist_ok=True)
    os.makedirs(os.path.dirname(os.path.abspath(a.output)), exist_ok=True)
    with open(spec_path, "w", encoding="utf-8") as fh:
        json.dump(spec, fh, ensure_ascii=False, indent=1)

    print(f"Cot file mau: {len(headers)} · map duoc: {len(headers) - len(unmapped)}")
    if unmapped:
        print("Cot trong file mau KHONG co du lieu (de trong, giu nguyen vi tri):")
        for h in unmapped:
            print(f"   - {h}")
    if unused:
        print("Du lieu goc KHONG co cot tuong ung trong file mau (field chua tao tren Base?):")
        for k in unused:
            print(f"   - {k}")
    sys.exit(subprocess.run([sys.executable, WRITER, "--spec", spec_path]).returncode)


if __name__ == "__main__":
    main()
